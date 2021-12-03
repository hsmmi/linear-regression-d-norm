import os
from random import randint
import numpy as np

def build_testcase(number_of_value,d_norm,lenght_of_value):

    with open(os.path.join(os.path.dirname(__file__),"dataset/mathAIH02.txt"), 'w') as f:
        f.write(f'd_norm in norm is: {d_norm}\n')
 
    tmp = [None] * number_of_value
    for i in range(number_of_value):
        tmp[i] = randint(1, lenght_of_value)-lenght_of_value/2
    tmp.sort()

    with open(os.path.join(os.path.dirname(__file__),"dataset/mathAIH02.txt"), 'a') as f:
        f.write(f'Values are: {tmp}\n')

    weighting, biasing = randint(-lenght_of_value, lenght_of_value), randint(-lenght_of_value, lenght_of_value)
    for i in range(number_of_value):
        tmp[i] = tmp[i]*weighting+biasing+ randint(1, lenght_of_value/5)-lenght_of_value/10
    # tmp.sort()

    with open(os.path.join(os.path.dirname(__file__),"dataset/mathAIH02.txt"), 'a') as f:
        f.write(f'Target values are: {tmp}\n')

    with open(os.path.join(os.path.dirname(__file__),"dataset/mathAIH02.txt"), 'a') as f:
        f.write(f'Weighting and biassing values are: [{weighting}, {biasing}]\n')

def read_testcase(file, printer = 0):
    with open(os.path.join(os.path.dirname(__file__),file),'r') as f:
        d_norm = int(f.readline().split(':')[1])
        if(d_norm > 2): d_norm = np.inf
        source_value = np.array(list(map(float, f.readline().split(':')[1][2:-2].split(', '))))
        target_value = np.array(list(map(float, f.readline().split(':')[1][2:-2].split(', '))))
        wb = np.array(list(map(float, f.readline().split(':')[1][2:-2].split(', '))))
    print(f'Weighting and biassing values are:\n{wb}\n')
    if(printer):
        print(f'd_norm in norm is:\n{d_norm}\n')
        print(f'Values are:\n{source_value}\n')
        print(f'Target values are:\n{target_value}\n')
    
    return source_value, target_value, d_norm

def read_dataset(file, atr):
    if (type(atr) == int):
        with open(os.path.join(os.path.dirname(__file__),file),'r') as f:
            return list(map(lambda x: float(x.split(',')[atr]), f.read().splitlines()))
    else:
        with open(os.path.join(os.path.dirname(__file__),file),'r') as f:
            return list(map(lambda x: [float(i) for i in (x.split(',')[atr[0]:atr[1]])], f.read().splitlines()))

def read_dataset_with_pandas(file, atr= None):
    import pandas as pd
    col_name = pd.read_csv(os.path.join(os.path.dirname(__file__),file),nrows=0).columns
    if (type(atr) == int):
        col_name = [col_name[atr]]
    elif(atr != None):
        col_name = col_name[atr[0]:atr[1]]
    data = pd.read_csv(os.path.join(os.path.dirname(__file__),file),usecols=col_name)

    return col_name, data

def read_dataset_with_pandas_to_nparray(file, atr= None):
    data = read_dataset_with_pandas(file, atr)[1]
    data = data.to_numpy()
    if (type(atr) == int):
        data = np.array(list(map(lambda x:x[0], data)))
    if(data.dtype == 'int'):
        data = data.astype('float')
    return data

def dataframe_to_docx_table(header,data,file,doc=None,save=1):
    """
    Read header and data
    If you gave if doc it add header and data to it and return it
    If you gave it save=0 it will not be save doc
    Return doc include header and data
    """
    import docx
    if(doc == None):
        doc = docx.d_normocument()
    doc.add_heading(header, 1)

    table = doc.add_table(rows=len(data.index)+1, cols=len(data.columns)+1)

    for j in range(len(data.columns)):
        table.cell(0,j+1).text = f'{data.columns[j]}'

    for i in range(len(data.index)):
        table.cell(i+1,0).text = f'{data.index[i]}'
        for j in range(len(data.columns)):
            table.cell(i+1,j+1).text = f'{data.iat[i,j]}'
    table.style = 'Table Grid'
    if(save):
        doc.save(file)
    return doc

def string_to_dataframe(string):
    from io import StringIO
    import pandas as pd
    data = StringIO(string)
    return pd.read_csv(data)

